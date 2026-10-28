from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from pptx import Presentation

from scripts.lyrics_format import editor_html_to_markdown, formatted_lines, lyrics_html
from scripts.markdown_to_pptx import build_slides, convert, plain_text
from scripts.print_song_sheet import _lyrics, create_print_song_sheet
from scripts.song_sheet import create_song_sheet
from scripts.visual_editor import merge_sections, split_sections, visual_song_editor


class VisualEditorTests(unittest.TestCase):
    def test_catalog_is_preserved_without_modifications(self):
        for path in (Path(__file__).resolve().parents[1] / "chants").rglob("*.md"):
            with self.subTest(path=path.name):
                text = path.read_text(encoding="utf-8")
                self.assertEqual(merge_sections(text, {}), text)

    def test_structure_stays_hidden_and_order_is_preserved(self):
        text = "<!-- ordre: C1 R C1 -->\n\n<!-- refrain -->\nRefrain\n\n---\n\n<!-- couplet: C1 -->\nCouplet\n"
        sections = split_sections(text)
        self.assertEqual([section.label for section in sections], ["Refrain", "Couplet 1"])
        self.assertEqual(sections[0].text, "Refrain")
        result = merge_sections(text, {"1": "<b>Paroles</b> et <u>suite</u>"})
        self.assertIn("<!-- ordre: C1 R C1 -->", result)
        self.assertIn("<!-- couplet: C1 -->", result)
        self.assertEqual(build_slides(result), ["**Paroles** et <u>suite</u>", "Refrain", "**Paroles** et <u>suite</u>"])

    def test_bold_and_underline_can_overlap_without_visible_syntax(self):
        markdown = editor_html_to_markdown("<div>Mon <b>chant <u>joyeux</u></b></div><div>Deuxième ligne</div>")
        self.assertEqual(plain_text(markdown), "Mon chant joyeux\nDeuxième ligne")
        html = lyrics_html(markdown)
        self.assertIn("<b><u>joyeux</u></b>", html)
        self.assertNotIn("**", html)
        self.assertIn("<br>Deuxième ligne", html)

    def test_text_and_paste_cannot_inject_html_or_song_directives(self):
        markdown = editor_html_to_markdown('<script>alert(1)</script><b onclick="bad()">OK</b><img src=x onerror="bad()">&lt;!-- refrain --&gt; * _ &amp;')
        self.assertNotIn("alert", markdown)
        self.assertNotIn("onclick", markdown)
        self.assertNotIn("<img", markdown)
        self.assertNotIn("<!--", markdown)
        html = lyrics_html(markdown)
        self.assertIn("&lt;!-- refrain --&gt;", html)
        self.assertIn("* _ &amp;", html)
        self.assertNotIn("<script", html)

    def test_newlines_and_leading_spaces_in_bold_are_valid_markdown(self):
        markdown = editor_html_to_markdown("<b> un chant <br> une suite </b>")
        self.assertEqual(plain_text(markdown), "un chant\n une suite")
        self.assertTrue(all(fragment.bold for line in formatted_lines(markdown) for fragment in line if fragment.text.strip()))

    def test_invalid_component_changes_are_rejected(self):
        for value in ({"8": "x"}, {"-1": "x"}, {"0": None}, []):
            with self.subTest(value=value), self.assertRaises(ValueError):
                merge_sections("Paroles", value)

    def test_markdown_syntax_typed_as_text_does_not_change_song_structure(self):
        text = "# Titre\n---\n[texte](lien)\n**étoiles**"
        markdown = editor_html_to_markdown(text)
        self.assertEqual(len(build_slides(markdown)), 1)
        self.assertEqual(plain_text(markdown), text)

    def test_save_event_is_processed_once_but_draft_remains_available(self):
        value = {"event_id": "once", "action": "save", "changes": {"0": "<u>Texte</u>"}}
        with patch("scripts.visual_editor._component", return_value=value), \
                patch("scripts.visual_editor.st.session_state", {}):
            self.assertEqual(visual_song_editor("Texte", key="test"), ("<u>Texte</u>", "save"))
            self.assertEqual(visual_song_editor("Texte", key="test"), ("<u>Texte</u>", ""))

    def test_html_pdf_and_powerpoint_keep_bold_and_underline(self):
        text = "**<u>Mot accentué</u>** et suite"
        with TemporaryDirectory() as directory:
            source = Path(directory) / "song.md"
            source.write_text(text, encoding="utf-8")
            songs = [("Chant", source)]
            html = create_song_sheet("Messe", songs).decode()
            self.assertIn("<strong><u>Mot accentué</u></strong>", html)
            self.assertIn("<b><u>Mot accentué</u></b>", _lyrics(text))
            pdf, pages = create_print_song_sheet("Messe", songs)
            self.assertTrue(pdf.startswith(b"%PDF"))
            self.assertGreaterEqual(pages, 1)
            target = Path(directory) / "song.pptx"
            convert(source, target)
            pptx = Presentation(BytesIO(target.read_bytes()))
            runs = pptx.slides[0].shapes[0].text_frame.paragraphs[0].runs
            self.assertEqual(runs[0].text, "Mot accentué")
            self.assertTrue(runs[0].font.bold)
            self.assertTrue(runs[0].font.underline)
