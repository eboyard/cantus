import unittest
from io import BytesIO

from scripts.markdown_to_pptx import BACKGROUND_COLOR, build_slides, convert, read_slides, read_index_slides
from pptx import Presentation
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from streamlit.testing.v1 import AppTest
from app import create_presentation


class RefrainTests(unittest.TestCase):
    def test_c_est_par_ta_grace_order_and_export(self):
        source = Path(__file__).resolve().parents[1] / "chants" / "c_est_par_ta_grâce.md"
        slides = read_slides(source)
        self.assertEqual(len(slides), 4)
        self.assertTrue(slides[0].startswith("Tout mon être"))
        self.assertTrue(slides[1].startswith("C’est par ta grâce"))
        self.assertTrue(slides[2].startswith("Tu connais mes craintes"))
        self.assertEqual(slides[1], slides[3])
        self.assertFalse(any("<!--" in slide for slide in slides))
        data, count = create_presentation([("C'est par ta grâce", source)])
        self.assertEqual(count, 4)
        presentation = Presentation(BytesIO(data))
        self.assertTrue(presentation.slides[0].shapes[0].text.startswith("Tout mon être"))
        self.assertTrue(presentation.slides[3].shapes[0].text.startswith("C’est par ta grâce"))

    def test_explicit_order_can_repeat_skip_and_reorder_parts(self):
        song = (
            "<!-- ordre: C2 R R C1 P1 R -->\n"
            "<!-- refrain -->\nR\n---\n"
            "<!-- couplet: C1 -->\nA\n---\n"
            "<!-- couplet: C2 -->\nB\n---\n"
            "<!-- couplet: C3 -->\nC\n---\n"
            "<!-- pont: P1 -->\nPont"
        )
        self.assertEqual(build_slides(song), ["B", "R", "R", "A", "Pont", "R"])

    def test_order_supports_implicit_numbers_and_windows_newlines(self):
        song = "\r\n<!-- Ordre: c2 r c1 -->\r\n<!-- Refrain -->\r\nR\r\n---\r\nA\r\n---\r\nB"
        self.assertEqual(build_slides(song), ["B", "R", "A"])

    def test_numbered_verses_preserve_default_alternation(self):
        song = "<!-- refrain -->\nR\n---\n<!-- couplet: C1 -->\nA\n---\n<!-- couplet: C2 -->\nB"
        self.assertEqual(build_slides(song), ["R", "A", "R", "B"])

    def test_invalid_order_is_rejected(self):
        songs = (
            "<!-- ordre: -->\nA",
            "<!-- ordre: R C2 -->\n<!-- refrain -->\nR\n---\nA",
            "<!-- ordre: C1 -->\n<!-- couplet: C1 -->\nA\n---\n<!-- couplet: C1 -->\nB",
            "<!-- ordre: P1 -->\n<!-- couplet: P1 -->\nA",
        )
        for song in songs:
            with self.subTest(song=song), self.assertRaises(ValueError):
                build_slides(song)

    def test_habiter_has_final_refrain(self):
        source = Path(__file__).resolve().parents[1] / "chants" / "habiter_le_seuil_de_ta_maison_seigneur.md"
        slides = read_slides(source)
        self.assertEqual(len(slides), 5)
        self.assertEqual([slides[i] for i in (0, 2, 4)], [slides[0]] * 3)
        self.assertTrue(slides[1].startswith("Veiller pour être prêt"))
        self.assertTrue(slides[3].startswith("Veiller pour accomplir"))
        self.assertFalse(any("<!--" in slide for slide in slides))

    def test_unmarked_song_keeps_original_order(self):
        self.assertEqual(build_slides("A\n---\nB\n---\nC"), ["A", "B", "C"])

    def test_refrain_between_verses(self):
        song = "<!-- refrain -->\nR\n---\nA\n---\nB\n---\nC"
        self.assertEqual(build_slides(song), ["R", "A", "R", "B", "R", "C"])

    def test_bridge_is_not_a_verse(self):
        song = "<!-- refrain -->\nR\n---\nA\n---\nB\n---\n<!-- pont -->\nP"
        self.assertEqual(build_slides(song), ["R", "A", "R", "B", "P"])

    def test_explicit_verse_markers_and_windows_newlines(self):
        song = "\r\n<!-- Refrain -->\r\nR\r\n --- \r\n<!-- couplet -->\r\nA\r\n---\r\nB\r\n---\r\n"
        self.assertEqual(build_slides(song), ["R", "A", "R", "B"])

    def test_refrain_already_between_verses_is_not_duplicated(self):
        self.assertEqual(build_slides("A\n---\n<!-- refrain -->\nR\n---\nB"), ["A", "R", "B"])

    def test_empty_or_multiple_refrains_are_rejected(self):
        for song in ("<!-- refrain -->", "<!-- refrain -->\nR\n---\n<!-- refrain -->\nS"):
            with self.subTest(song=song), self.assertRaises(ValueError):
                build_slides(song)

    def test_saint_est_son_nom_has_four_verses_and_bridge(self):
        source = Path(__file__).resolve().parents[1] / "chants" / "saint_est_son_nom.md"
        slides = read_slides(source)
        self.assertEqual(len(slides), 9)
        self.assertEqual([slides[i] for i in (0, 2, 4, 6)], [slides[0]] * 4)
        self.assertTrue(slides[8].startswith("Crions de joie !"))
        self.assertFalse(any("<!--" in slide for slide in slides))


class IndexTests(unittest.TestCase):
    def test_index_order_refrains_and_missing_lyrics(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "chant.md").write_text("<!-- refrain -->\nR\n---\nA\n---\nB", encoding="utf-8")
            (root / "fin.md").write_text("Fin", encoding="utf-8")
            index = root / "index.md"
            index.write_text(
                "| Ordre | Moment | Titre | Fichier |\n"
                "| --- | --- | --- | --- |\n"
                "| 1 | Entrée | Chant | [Markdown](chant.md) |\n"
                "| 2 | Psaume | Psaume 22 (23) | Aucun fichier existant |\n"
                "| 3 | Envoi | Fin | [Markdown](fin.md) |\n",
                encoding="utf-8",
            )
            self.assertEqual(read_index_slides(index), ["R", "A", "R", "B", "", "Psaume 22 (23)", "", "Fin"])
            destination = root / "presentation.pptx"
            self.assertEqual(convert(index, destination, from_index=True), 8)
            presentation = Presentation(destination)
            self.assertEqual(len(presentation.slides), 8)
            for position in (4, 6):
                slide = presentation.slides[position]
                self.assertEqual(len(slide.shapes), 0)
                self.assertEqual(slide.background.fill.fore_color.rgb, BACKGROUND_COLOR)

    def test_broken_link_is_an_error(self):
        with TemporaryDirectory() as directory:
            index = Path(directory) / "index.md"
            index.write_text("| 1 | Entrée | Chant | [Markdown](absent.md) |", encoding="utf-8")
            with self.assertRaises(FileNotFoundError):
                read_index_slides(index)

    def test_empty_index_is_an_error(self):
        with TemporaryDirectory() as directory:
            index = Path(directory) / "index.md"
            index.write_text("# Index sans chants", encoding="utf-8")
            with self.assertRaises(ValueError):
                read_index_slides(index)


class WebAppTests(unittest.TestCase):
    def test_preview_error_does_not_interrupt_app(self):
        app = AppTest.from_string("from app import main\nmain()")
        app.session_state["song_order"] = ["chants/c_est_par_ta_grâce.md"]
        with patch("app.read_slides", side_effect=ValueError("Partie C1 absente")):
            app.run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.error), 1)
        self.assertIn("c_est_par_ta_grâce.md", app.error[0].value)
        self.assertIn("Partie C1 absente", app.error[0].value)
        self.assertTrue(any(button.label == "Créer le PowerPoint" for button in app.button))

    def test_create_presentation_assembles_songs_in_order(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            first = root / "first.md"
            second = root / "second.md"
            first.write_text("Chant d'entrée", encoding="utf-8")
            second.write_text("Chant d'envoi", encoding="utf-8")

            data, slide_count = create_presentation([("Entrée", first), ("Envoi", second)])

            presentation = Presentation(BytesIO(data))
            self.assertEqual(slide_count, 3)
            self.assertIn("Chant d'entrée", presentation.slides[0].shapes[0].text)
            self.assertIn("Chant d'envoi", presentation.slides[2].shapes[0].text)


if __name__ == "__main__":
    unittest.main()
