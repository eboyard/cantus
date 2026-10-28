"""Convertit les paroles Markdown en diapositives PowerPoint au format 16:9."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
import re
from urllib.parse import unquote

if __package__:
    from scripts.lyrics_format import formatted_lines
else:
    from lyrics_format import formatted_lines

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


# Style relevé dans diaporamas/2026-rameaux.pptx : fond du masque
# lt2 (#000080), avec une luminosité de 50 %, soit #000040.
BACKGROUND_COLOR = RGBColor(0, 0, 64)
FONT_NAME = "Comic Sans MS"
TITLE_TAG = re.compile(r"^[ \t]*<!--[ \t]*titre[ \t]*:([^\r\n]*?)-->[ \t]*\r?$", re.MULTILINE | re.IGNORECASE)
METADATA_TAG = re.compile(
    r"^[ \t]*<!--[ \t]*(auteur|accords|tonalite|informations)[ \t]*:([^\r\n]*?)-->[ \t]*\r?$",
    re.MULTILINE | re.IGNORECASE,
)


def read_title(path: Path) -> str:
    """Utilise le titre déclaré, sinon le titre déduit du nom du fichier."""
    marker = TITLE_TAG.search(path.read_text(encoding="utf-8-sig"))
    if marker and marker.group(1).strip():
        return marker.group(1).strip()
    return path.stem.replace("_", " ").replace("-", " ").title()


def read_slides(path: Path) -> list[str]:
    """Lit le chant selon son ordre explicite ou l'alternance par défaut."""
    return build_slides(path.read_text(encoding="utf-8-sig"))


def read_index_slides(path: Path) -> list[str]:
    """Assemble les chants liés dans un index, dans l'ordre de son tableau."""
    slides: list[str] = []
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) != 4 or not cells[0].isdigit():
            continue
        if slides:
            slides.append("")
        link = re.search(r"\[[^\]]*\]\(([^)]+)\)", cells[3])
        if link:
            song_path = path.parent / unquote(link.group(1))
            song_slides = read_slides(song_path)
            if not song_slides:
                raise ValueError(f"Le chant {song_path} ne contient aucune parole.")
            slides.extend(song_slides)
        else:
            # Conserve la place d'une entrée dont les paroles ne sont pas disponibles.
            slides.append(cells[2])
    if not slides:
        raise ValueError("L'index ne contient aucune entrée dans son tableau.")
    return slides


@dataclass(frozen=True)
class SongPart:
    identifier: str
    kind: str
    text: str


def build_slides(markdown: str) -> list[str]:
    return [part.text for part in build_song_parts(markdown)]


def build_song_parts(markdown: str) -> list[SongPart]:
    """Les commentaires de type sont placés au début d'un bloc séparé par ---.

    La directive ordre en tête de fichier choisit les parties à projeter.
    Sans directive, un bloc sans marque est un couplet et le pont interrompt
    la succession de couplets ; aucun refrain n'est ajouté autour du pont.
    """
    markdown = METADATA_TAG.sub("", TITLE_TAG.sub("", markdown))
    order_marker = re.match(r"\A\s*<!--[ \t]*ordre[ \t]*:([^\r\n]*?)-->[ \t]*(?:\r?\n|$)", markdown, re.IGNORECASE)
    order = None
    if order_marker:
        order = order_marker.group(1).upper().split()
        if not order:
            raise ValueError("La directive ordre ne contient aucune partie.")
        markdown = markdown[order_marker.end():]

    sections = re.split(r"(?m)^[ \t]*---[ \t]*\r?$", markdown)
    blocks: list[SongPart] = []
    parts: dict[str, SongPart] = {}
    counts = {"couplet": 0, "pont": 0}
    for section in sections:
        text = section.strip()
        if not text:
            continue
        marker = re.match(r"\A<!--[ \t]*(refrain|couplet|pont)(?:[ \t]*:[ \t]*([CP][1-9][0-9]*))?[ \t]*-->[ \t]*(?:\r?\n|$)", text, re.IGNORECASE)
        kind = marker.group(1).lower() if marker else "couplet"
        if marker:
            text = text[marker.end():].strip()
            if not text:
                raise ValueError(f"Le bloc marqué {kind} ne contient aucune parole.")
        if kind == "refrain":
            identifier = "R"
            if marker and marker.group(2):
                raise ValueError("Le refrain utilise l'identifiant R, sans suffixe.")
        else:
            counts[kind] += 1
            prefix = "C" if kind == "couplet" else "P"
            identifier = marker.group(2).upper() if marker and marker.group(2) else f"{prefix}{counts[kind]}"
            if not identifier.startswith(prefix):
                raise ValueError(f"L'identifiant {identifier} ne correspond pas au type {kind}.")
        if order is not None and identifier in parts:
            raise ValueError(f"L'identifiant {identifier} est défini plusieurs fois.")
        part = SongPart(identifier, kind, text)
        parts[identifier] = part
        blocks.append(part)

    refrains = [part for part in blocks if part.kind == "refrain"]
    if len(refrains) > 1:
        raise ValueError("Un chant ne peut contenir qu'un bloc marqué refrain.")

    if order is not None:
        for identifier in order:
            if identifier not in parts:
                raise ValueError(
                    f"La partie {identifier} indiquée dans l'ordre n'existe pas. "
                    "Vérifiez son identifiant et séparez les blocs par une ligne contenant uniquement '---'."
                )
        return [parts[identifier] for identifier in order]

    slides: list[SongPart] = []
    previous_kind = None
    for part in blocks:
        if refrains and part.kind == previous_kind == "couplet":
            slides.append(refrains[0])
        slides.append(part)
        previous_kind = part.kind
    return slides


def styled_lines(section: str) -> list[list[tuple[str, bool]]]:
    """Compatibilité des lecteurs de texte et de gras existants."""
    return [[(fragment.text, fragment.bold) for fragment in line]
            for line in formatted_lines(section)]


def plain_text(section: str) -> str:
    """Retire la mise en forme Markdown des paroles."""
    return "\n".join("".join(text for text, _ in line) for line in styled_lines(section))


def fit_font(text: str, maximum: int) -> int:
    """Estime l'espace occupé, y compris les lignes susceptibles d'être repliées."""
    # Zone de texte : 11,73 × 6,3 pouces ; marge prudente pour les caractères larges.
    for size in range(maximum, 0, -1):
        characters_per_line = max(1, int(11.73 * 72 / (size * 0.65)))
        rows = sum(
            max(1, (len(line) + characters_per_line - 1) // characters_per_line)
            for line in text.splitlines()
        )
        if rows * size * 1.25 <= 6.3 * 72:
            return size
    return 1


def convert(source: Path, destination: Path, font_size: int = 40, *, from_index: bool = False) -> int:
    sections = read_index_slides(source) if from_index else read_slides(source)
    if not sections:
        raise ValueError("Le fichier Markdown ne contient aucune parole.")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)

    for section in sections:
        lines = formatted_lines(section)
        text = "\n".join("".join(fragment.text for fragment in line) for line in lines)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = BACKGROUND_COLOR
        if not text:
            continue
        box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.73), Inches(6.3))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        size = fit_font(text, font_size)

        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            for content, bold, underline in line:
                run = paragraph.add_run()
                run.text = content
                run.font.bold = bold
                run.font.underline = underline
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.15
            paragraph.font.name = FONT_NAME
            paragraph.font.size = Pt(size)
            paragraph.font.color.rgb = RGBColor(255, 255, 255)

    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(destination))
    return len(sections)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Fichier Markdown encodé en UTF-8")
    output = parser.add_mutually_exclusive_group()
    output.add_argument("-o", "--output", type=Path, help="Fichier .pptx à créer")
    output.add_argument("--output-dir", type=Path, help="Dossier de sortie ; le nom est déduit du fichier source")
    parser.add_argument("--index", action="store_true", help="Assembler les chants liés dans le tableau d'un index Markdown")
    parser.add_argument("--font-size", type=int, default=40, help="Taille maximale en points (40 par défaut)")
    args = parser.parse_args()
    if not 1 <= args.font_size <= 200:
        parser.error("--font-size doit être compris entre 1 et 200.")
    destination = args.output or args.source.with_suffix(".pptx")
    if args.output_dir is not None:
        name = "messe_" + args.source.stem.removeprefix("index_") if args.index else args.source.stem
        destination = args.output_dir / f"{name}.pptx"
    if destination.suffix.lower() != ".pptx":
        parser.error("Le fichier de sortie doit avoir l'extension .pptx.")
    try:
        count = convert(args.source, destination, args.font_size, from_index=args.index)
    except (OSError, ValueError) as exc:
        parser.exit(1, f"Erreur : {exc}\n")
    print(f"Présentation créée : {destination} ({count} diapositives)")


if __name__ == "__main__":
    main()
